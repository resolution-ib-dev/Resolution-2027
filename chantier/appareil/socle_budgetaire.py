# -*- coding: utf-8 -*-
"""Socle budgétaire — le PLF et son interprétation, lus dans les classeurs.

Les classeurs de l'auteur restent la **source officielle** : ils ne se
réécrivent pas, ils ne viennent pas au dépôt, et rien ici ne prétend les
remplacer. Ce module en produit une **lecture formalisée**, à trois fins.

1. **Rendre la chaîne lisible.** Un agrégat du `REF_doctrine` — 10,6 Md€ pour
   France Compétences — se recompose ici jusqu'aux lignes de taxe affectée qui
   le portent, avec leur SIREN et leur référence juridique.
2. **Prouver la grille de lecture.** Ce que ce module recompose doit retrouver
   ce que le classeur affiche. Tant que le bouclage n'est pas vérifié, la grille
   est une hypothèse ; une fois vérifié, elle est un fait.
3. **Permettre le rejeu.** Les nomenclatures sont stables d'un exercice à
   l'autre : SIREN, numéro de dépense fiscale, numéro de programme. Un PLF neuf
   se lit avec le même module, et la comparaison des deux socles dit ce qui a
   bougé.

## Les deux couches, et pourquoi elles ne se mélangent pas

**Le socle** est ce que le PLF publie : un bénéficiaire, un montant, une
référence. Il ne se discute pas.

**L'interprétation** est ce que l'auteur a ajouté en colonnes : le régime retenu
— supprimer, fusionner, laisser au flux outre-mer — et les montants qui en
découlent. Elle se discute, elle se varie, et c'est elle qu'un PLF neuf oblige à
réexaminer.

Les deux sont dans le même classeur, côte à côte. Ici elles sont dans deux blocs
distincts de chaque entrée, `socle` et `interpretation`, de sorte qu'on puisse
changer la seconde sans toucher au premier.

## Ce que ce module ne fait pas

Il n'invente aucun régime, aucun taux, aucune affectation. Une cellule vide
reste vide : le classeur ne dit rien, le socle ne dit rien. Il ne corrige pas
davantage un montant qui paraîtrait faux — ce serait substituer une source à
l'autre, et l'auteur garde la sienne.

Usage : python3 socle_budgetaire.py <sortie.json> \\
            --taxes taxes_affectees.xlsx \\
            --operateurs synthese_etp.xlsx \\
            --depenses-fiscales depenses_fiscales.xlsx \\
            --depenses-bg depenses_bg.xlsx \\
            --calculs synthese_calculs.xlsx

Le classeur de calculs fait exception à la règle des deux couches : il est
entièrement de l'auteur, socle et interprétation confondus. Ses entrées portent
donc `classeur` et `lecture`, et non `socle` et `interpretation`.

Les classeurs sont des pièces jointes du projet, en `.xls`. Ils se convertissent
une fois dans l'atelier :

    soffice --headless --convert-to xlsx --outdir /tmp <fichier>.xls
"""
import json
import sys

import nomenclature_lolf as nl

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None


def _txt(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(' ', '').replace('\xa0', '').replace(' ', '')
    s = s.replace(',', '.')
    try:
        return float(s)
    except ValueError:
        return None


# --------------------------------------------------------- les taxes affectées
# Annexe 2 du tome I des Voies et moyens, une ligne par taxe affectée. Les
# en-têtes tiennent sur deux lignes, 15 et 16 ; les données commencent en 17.
#
# Colonnes du socle : le secteur et la nature juridique du bénéficiaire, son
# SIREN, sa mission et son programme de rattachement, le code et le libellé de
# la taxe, sa référence juridique, et l'affectation nette.
#
# Colonnes de l'interprétation, ajoutées par l'auteur : le régime de suppression,
# et les deux temps de la restitution.
#
# **Le vocabulaire compte, et il n'est pas le même qu'en dépense fiscale.**
# En budgétaire — taxes affectées et crédits — le « gage CSG » est ce qui est
# **restituable dès l'année 1**, et le « solde » ce qui est **restitué ensuite**.
# La vraie économie valorisable restituée est **leur total**. Les nommer
# « gage » et « économie en sus » laissait croire à deux natures différentes ;
# ce sont deux temps de la même restitution.
TAXES_DEBUT = 17
TAXES_COLONNES = {
    'secteur': 'A', 'categorie': 'B', 'nature_juridique': 'C', 'siren': 'D',
    'programme': 'E', 'mission': 'F', 'programme_libelle': 'G',
    'code_taxe': 'H', 'beneficiaire': 'I',
    'regime': 'J', 'gage_csg': 'K', 'economie': 'L', 'taxe_affectee': 'M',
    'affectation_nette': 'O', 'reference_juridique': 'Y',
}


def des_taxes(chemin):
    ws = openpyxl.load_workbook(chemin, data_only=True)['taxes affectées']
    out = []
    for i in range(TAXES_DEBUT, ws.max_row + 1):
        v = {k: ws[f'{c}{i}'].value for k, c in TAXES_COLONNES.items()}
        if not _txt(v['beneficiaire']):
            continue
        out.append({
            'id': f"TA-{_txt(v['siren']) or '000000000'}-{_txt(v['code_taxe'])}"
                  f"-{i}",
            'ligne': i,
            'socle': {
                'secteur': _txt(v['secteur']),
                'categorie': _txt(v['categorie']),
                'nature_juridique': _txt(v['nature_juridique']),
                'siren': _txt(v['siren']),
                'mission': _txt(v['mission']),
                'programme': _txt(v['programme']),
                'programme_libelle': _txt(v['programme_libelle']),
                'code_taxe': _txt(v['code_taxe']),
                'beneficiaire': _txt(v['beneficiaire']),
                'reference_juridique': _txt(v['reference_juridique']),
                'affectation_nette_eur': _num(v['affectation_nette']),
            },
            'interpretation': {
                'regime': _txt(v['regime']),
                'restitue_annee_1_eur': _num(v['gage_csg']),
                'restitue_ensuite_eur': _num(v['economie']),
                'economie_restituee_totale_eur': (
                    (_num(v['gage_csg']) or 0.0)
                    + (_num(v['economie']) or 0.0)
                    if (_num(v['gage_csg']) is not None
                        or _num(v['economie']) is not None) else None),
                'taxe_affectee_m_eur': _num(v['taxe_affectee']),
            },
        })
    return out


# ------------------------------------------------------------- les opérateurs
# Onglet « Opérateurs » du classeur ETP et agences. Une ligne par opérateur du
# PLF, avec ses emplois sous et hors plafond aux deux exercices, et le régime
# retenu marqué d'une croix dans l'une des cinq colonnes N à R.
#
# **Un opérateur porte au plus un régime.** Le contrôle le vérifie : une ligne à
# deux croix serait comptée deux fois dans les agrégats du classeur.
OPERATEURS_DEBUT = 5
OPERATEURS_COLONNES = {
    'mission_code': 'A', 'mission': 'B', 'programme': 'C',
    'programme_libelle': 'D', 'operateur': 'E', 'nombre': 'F',
    'etpt_total_lfi': 'G', 'etpt_plafond_lfi': 'H', 'etpt_hors_lfi': 'I',
    'etpt_total_plf': 'J', 'etpt_plafond_plf': 'K', 'etpt_hors_plf': 'L',
}
OPERATEURS_REGIMES = {
    'N': 'vente', 'O': 'suppression', 'P': 'epic_musee',
    'Q': 'epic_enseignement', 'R': 'internalisation',
}


def des_operateurs(chemin):
    ws = openpyxl.load_workbook(chemin, data_only=True)['Opérateurs']
    out = []
    for i in range(OPERATEURS_DEBUT, ws.max_row + 1):
        v = {k: ws[f'{c}{i}'].value for k, c in OPERATEURS_COLONNES.items()}
        nom = _txt(v['operateur'])
        if not nom or nom.lower().startswith('total'):
            continue
        regimes = [r for col, r in OPERATEURS_REGIMES.items()
                   if _txt(ws[f'{col}{i}'].value)]
        out.append({
            'id': f"OP-{i}",
            'ligne': i,
            'socle': {
                'operateur': nom,
                'mission': _txt(v['mission']),
                'programme': _txt(v['programme']),
                'programme_libelle': _txt(v['programme_libelle']),
                'nombre_entites': _num(v['nombre']),
                'etpt_total_lfi_2025': _num(v['etpt_total_lfi']),
                'etpt_sous_plafond_lfi_2025': _num(v['etpt_plafond_lfi']),
                'etpt_hors_plafond_lfi_2025': _num(v['etpt_hors_lfi']),
                'etpt_total_plf_2026': _num(v['etpt_total_plf']),
                'etpt_sous_plafond_plf_2026': _num(v['etpt_plafond_plf']),
                'etpt_hors_plafond_plf_2026': _num(v['etpt_hors_plf']),
            },
            'interpretation': {
                'regime': regimes[0] if len(regimes) == 1 else None,
                'regimes_multiples': regimes if len(regimes) > 1 else [],
                'sans_regime': not regimes,
            },
        })
    return out


# -------------------------------------------------------- les dépenses fiscales
# Onglet « Chiffrages IB » de l'annexe 3 du tome II. Les en-têtes tiennent sur
# les lignes 19 et 20, les données commencent en 21. Le numéro de dépense
# fiscale est la clé du PLF et il est stable d'un exercice à l'autre.
#
# La colonne I porte la réalisation 2024 telle que le PLF la publie ; la colonne
# J la reprend augmentée de la TVA des administrations publiques, retraitement
# de l'auteur. Les deux se gardent : la première est du socle, la seconde de
# l'interprétation.
DF_DEBUT = 21
DF_COLONNES = {
    'categorie': 'A', 'sous_categorie': 'B', 'sous_sous_categorie': 'C',
    'numero': 'D', 'libelle': 'E', 'creation': 'F', 'fin_fait_generateur': 'G',
    'beneficiaires': 'H', 'realisation_2024': 'I',
    'realisation_yc_tva': 'J', 'part_economie_nette': 'K',
    'part_effet_macro': 'L', 'regime': 'M', 'gage_net_csg_m_eur': 'N',
    'effet_macro_m_eur': 'O', 'entreprise': 'P', 'nature_niche': 'Q',
}


def des_depenses_fiscales(chemin):
    ws = openpyxl.load_workbook(chemin, data_only=True)['Chiffrages IB']
    out = []
    for i in range(DF_DEBUT, ws.max_row + 1):
        v = {k: ws[f'{c}{i}'].value for k, c in DF_COLONNES.items()}
        num = _txt(v['numero'])
        if not num:
            continue
        out.append({
            'id': f"DF-{num}",
            'ligne': i,
            'socle': {
                'numero': num,
                'categorie': _txt(v['categorie']),
                'sous_categorie': _txt(v['sous_categorie']),
                'sous_sous_categorie': _txt(v['sous_sous_categorie']),
                'libelle': _txt(v['libelle']),
                'creation': _txt(v['creation']),
                'fin_fait_generateur': _txt(v['fin_fait_generateur']),
                'beneficiaires': _num(v['beneficiaires']),
                # « ε » et « nc » sont des mentions du PLF, non des nombres :
                # elles se gardent telles quelles et ne se lisent pas comme zéro.
                'realisation_2024_m_eur': _num(v['realisation_2024']),
                'realisation_2024_mention': (
                    _txt(v['realisation_2024'])
                    if _num(v['realisation_2024']) is None else None),
            },
            'interpretation': {
                'realisation_yc_tva_apu_m_eur': _num(v['realisation_yc_tva']),
                'part_economie_nette': _num(v['part_economie_nette']),
                'part_effet_macro': _num(v['part_effet_macro']),
                'regime': _txt(v['regime']),
                'gage_net_csg_m_eur': _num(v['gage_net_csg_m_eur']),
                'effet_macro_m_eur': _num(v['effet_macro_m_eur']),
                'entreprise': _txt(v['entreprise']),
                'nature_niche': _txt(v['nature_niche']),
            },
        })
    return out


# ------------------------------------------------------------- la synthèse agences
# Onglet « Synthèse agences ». Quatre familles d'entités, six régimes, et un
# total de 1 104 agences d'État. C'est le tableau que le REF_doctrine cite.
AGENCES_LIGNES = {5: 'agences_etat', 6: 'operateurs_plf',
                  7: 'organismes_hors_plf', 8: 'autorites_independantes',
                  9: 'commissions', 10: 'agences_locales'}
AGENCES_COLONNES = {'D': 'total', 'E': 'cession_actif', 'F': 'suppression',
                    'G': 'internalisation', 'H': 'epic_musee',
                    'I': 'epic_enseignement', 'J': 'autorite_independante'}


def des_agences(chemin):
    ws = openpyxl.load_workbook(chemin, data_only=True)['Synthèse agences']
    out = {}
    for i, nom in AGENCES_LIGNES.items():
        out[nom] = {v: _num(ws[f'{c}{i}'].value)
                    for c, v in AGENCES_COLONNES.items()}
    return out


# --------------------------------------------- les dépenses du budget général
# Onglet « Données PAP 2026 » du classeur des dépenses. Une ligne par
# croisement mission × programme × action × sous-action × catégorie de dépense.
# C'est la maille la plus fine du budget de l'État, et **elle est intégralement
# du socle** : l'interprétation, elle, est agrégée à l'onglet « Synthèse ».
#
# Les catégories de dépense suivent la nomenclature par nature de la LOLF :
# 21 rémunérations, 22 cotisations, 23 prestations sociales, 31 fonctionnement,
# 32 subventions pour charges de service public — les opérateurs —, 5
# investissement, 61 à 64 transferts aux ménages, entreprises, collectivités et
# autres.
PAP_COLONNES = {
    'type_mission': 'A', 'mission': 'B', 'code_mission': 'C',
    'programme': 'D', 'programme_libelle': 'E', 'action': 'F',
    'action_libelle': 'G', 'sous_action': 'H', 'sous_action_libelle': 'I',
    'categorie': 'J', 'titre': 'K', 'ae_plf': 'L', 'cp_plf': 'M',
    'ae_fdc_adp': 'N', 'cp_fdc_adp': 'O',
}


def du_pap(chemin):
    ws = openpyxl.load_workbook(chemin, data_only=True)['Données PAP 2026']
    out = []
    for i in range(2, ws.max_row + 1):
        v = {k: ws[f'{c}{i}'].value for k, c in PAP_COLONNES.items()}
        if not _txt(v['programme']):
            continue
        out.append({
            'id': f"PAP-{_txt(v['programme'])}-{_txt(v['action']) or '00'}"
                  f"-{_txt(v['sous_action']) or '00'}-{_txt(v['categorie'])}",
            'ligne': i,
            'socle': {k: (_num(v[k]) if k.startswith(('ae_', 'cp_'))
                          else _txt(v[k])) for k in PAP_COLONNES},
        })
    return out


# L'onglet « Synthèse » du même classeur porte l'interprétation, agrégée par
# nature de dépense et par destinataire. Les cellules en erreur du classeur — le
# tableur y laisse des `#VALUE!` — se gardent telles quelles : **ce n'est pas au
# socle de réparer une formule cassée**, c'est à l'auteur de la voir.
BG_REPERES = {
    'salaires_m_eur': 'E4', 'cotisations_m_eur': 'F4',
    'prestations_m_eur': 'G4', 'fonctionnement_m_eur': 'H4',
    'investissement_m_eur': 'I4', 'operateurs_m_eur': 'K4',
    'transferts_menages_m_eur': 'M4', 'transferts_entreprises_m_eur': 'O4',
    'transferts_collectivites_m_eur': 'Q4', 'transferts_autres_m_eur': 'S4',
    'salaire_moyen_eur': 'E5', 'etp_etat_supprimes': 'E8',
    'salaires_non_regalien_m_eur': 'E6',
    'fonctionnement_non_regalien_m_eur': 'H6',
    'investissement_non_regalien_m_eur': 'I6',
}


# La grille du budget général — dix catégories, deux blocs de lignes, et une
# couche de décision qu'aucun autre document ne porte.
#
#   lignes 1 à 4     la grille : catégorie LOLF en tête, total des crédits de
#                    paiement, et pour les cinq catégories traitées la part
#                    supprimable dès l'année 1
#   lignes 5 à 23    les postes, en paires libellé / valeur, colonne par
#                    colonne. Le libellé est le seul endroit du corpus qui
#                    dise ce que le montant est.
#   lignes 25 à 59   les 35 missions, montant par catégorie
#   lignes 64 à fin  les 128 programmes, montant par catégorie **et le
#                    traitement retenu**, un mot par case
#
# **Le traitement est la couche qui manquait.** Un programme porte, pour chaque
# catégorie de transfert, un mot — « Oui », « En 3 ans », « Fusion CI »,
# « Bourse », « Sécu », « Flux OM » — et ce mot dit si le crédit est supprimé,
# reporté ou seulement déplacé. Sans lui, l'agrégat est un total ; avec lui,
# c'est un chiffrage.
BG_GRILLE = 4
BG_POSTES = (5, 23)
BG_MISSIONS = (25, 59)
BG_PROGRAMMES_DEBUT = 64
BG_REGALIEN = 'J'
# Les paramètres hors transferts : un intitulé en D, des valeurs par catégorie.
# (ligne du libellé, ligne de la valeur, catégories portant la valeur). Deux
# paramètres portent leur libellé une ligne au-dessus de leur montant : les lire
# sur la même ligne rendrait le nombre comme intitulé.
BG_PARAMETRES = (
    ('Salaire moyen', 5, 5, ('21',)),
    ('Assiette non régalienne', 6, 6, ('21', '31', '5')),
    ("Économie restituée en année 1", 7, 7, ('21', '31', '5')),
    ("ETP d'État supprimés", 8, 8, ('21',)),
    ("Fusion du crédit d'impôt, budgétaire et fiscal", 15, 16, ()),
    ('Bourse, crédits budgétaires et taxes affectées', 18, 19, ()),
)


def _pose(ws, colonne, ligne):
    return _num(ws[f'{colonne}{ligne}'].value)


def du_bg_synthese(chemin):
    ws = openpyxl.load_workbook(chemin, data_only=True)['Synthèse']
    out = {}
    # Les repères historiques restent en tête : d'autres modules les lisent.
    for nom, cel in BG_REPERES.items():
        v = ws[cel].value
        n = _num(v)
        out[nom] = n if n is not None else _txt(v)

    # --- la grille
    grille = []
    for code, c in sorted(nl.CATEGORIES.items(),
                          key=lambda kv: kv[1]['colonne']):
        grille.append({
            'categorie': code,
            'titre': c['titre'],
            'titre_libelle': nl.TITRES[c['titre']],
            'libelle_court': c['court'],
            'libelle_officiel': c['officiel'],
            'traitee_ligne_a_ligne': code in nl.CATEGORIES_TRAITEES,
            'credit_plf_m_eur': _pose(ws, c['colonne'], BG_GRILLE),
            'suppression_immediate_m_eur': (
                _pose(ws, c['colonne_traitement'], BG_GRILLE)
                if c.get('colonne_traitement') else None),
        })

    # --- les paramètres, intitulé en D, valeurs par catégorie
    parametres = []
    for intitule, i_lib, i_val, cats in BG_PARAMETRES:
        brut = _txt(ws[f'D{i_lib}'].value)
        for code in (cats or ('',)):
            col = nl.CATEGORIES[code]['colonne'] if code else 'D'
            val = _pose(ws, col, i_val)
            if val is None:
                continue
            q, cert = nl.qualifier(brut)
            parametres.append({
                'intitule': intitule,
                'libelle_classeur': brut,
                'categorie': code or None,
                'valeur': val,
                'qualification': q,
                'certitude': cert,
                'unite': nl.unite_de(q),
                'ligne': i_val,
            })

    # --- les postes, en paires libellé / valeur
    postes = []
    for code, c in nl.CATEGORIES.items():
        tr = c.get('colonne_traitement')
        if not tr:
            continue
        for i in range(BG_POSTES[0], BG_POSTES[1] + 1):
            lib = _txt(ws[f"{c['colonne']}{i}"].value)
            val = _pose(ws, tr, i)
            if not lib or val is None:
                continue
            q, cert = nl.qualifier(lib)
            postes.append({
                'categorie': code,
                'titre': c['titre'],
                'libelle': lib.split(' http')[0].strip(),
                'libelle_brut': lib,
                'valeur_m_eur': val,
                'qualification': q,
                'certitude': cert,
                'unite': nl.unite_de(q),
                'ligne': i,
            })

    # --- les missions
    missions = []
    for i in range(BG_MISSIONS[0], BG_MISSIONS[1] + 1):
        code = _txt(ws[f'B{i}'].value)
        if not code:
            continue
        missions.append({
            'code': code, 'ligne': i,
            'credit_plf_m_eur': {k: _pose(ws, c['colonne'], i)
                                 for k, c in nl.CATEGORIES.items()},
            'suppression_immediate_m_eur': {
                k: _pose(ws, c['colonne_traitement'], i)
                for k, c in nl.CATEGORIES.items()
                if c.get('colonne_traitement')},
        })

    # --- les programmes, avec leur traitement par catégorie
    programmes = []
    inconnus = set()
    for i in range(BG_PROGRAMMES_DEBUT, ws.max_row + 1):
        genre = _txt(ws[f'A{i}'].value)
        if genre != 'P':
            continue
        traitements = {}
        for k, c in nl.CATEGORIES.items():
            tr = c.get('colonne_traitement')
            if not tr:
                continue
            mot = _txt(ws[f'{tr}{i}'].value)
            if mot is None:
                continue
            if mot not in nl.TRAITEMENTS:
                inconnus.add(mot)
            traitements[k] = mot
        regalien = _txt(ws[f'{BG_REGALIEN}{i}'].value)
        programmes.append({
            'mission': _txt(ws[f'B{i}'].value),
            'programme': _txt(ws[f'C{i}'].value),
            'ligne': i,
            'regalien': regalien == 'X',
            'credit_plf_m_eur': {k: _pose(ws, c['colonne'], i)
                                 for k, c in nl.CATEGORIES.items()},
            'traitement': traitements,
        })

    out['grille'] = grille
    out['parametres'] = parametres
    out['postes'] = postes
    out['missions'] = missions
    out['programmes'] = programmes
    out['traitements_inconnus'] = sorted(inconnus)
    return out


# ------------------------------------- les dépenses fiscales, feuille par feuille
# L'annexe 3 tient neuf feuilles, toutes indexées par le **numéro de dépense
# fiscale**. Elles ne se lisent donc pas séparément : elles enrichissent la même
# entrée. Ce qu'elles ajoutent est du socle — c'est le PLF qui l'écrit.
#
#   Chiffrages                  réalisation 2024, prévisions 2025 et 2026
#   Echéances                   création, dernière modification, fin du fait
#                               générateur, fin d'incidence, suppression
#   Bénéficiaires               nature et nombre
#   Méthodologie                **fiabilité du chiffrage** et méthode retenue
#   Références juridiques       norme de référence, code et article
#   Programme de rattachement   programme budgétaire porteur
#
# La fiabilité est celle que l'administration déclare — « Très bonne », « Bonne »,
# « Ordre de grandeur » — et **elle qualifie chaque chiffre qu'on reprend**. Un
# gage bâti sur un ordre de grandeur ne vaut pas un gage bâti sur une simulation.
DF_FEUILLES = {
    'Chiffrages': (5, 'D', {'realisation_2024': 'F', 'prevision_2025': 'G',
                            'prevision_2026': 'H'}),
    'Echéances': (5, 'D', {'creation': 'F', 'derniere_modification': 'G',
                           'fin_fait_generateur': 'H',
                           'fin_incidence_budgetaire': 'I',
                           'suppression': 'J',
                           'motif_creation_suppression': 'L'}),
    'Bénéficiaires': (3, 'D', {'beneficiaire_nature': 'F',
                               'beneficiaire_nombre': 'G'}),
    'Méthodologie': (3, 'D', {'fiabilite_du_realise': 'F',
                              'methode_de_chiffrage': 'G',
                              'changement_de_methode': 'H'}),
    'Références juridiques': (3, 'D', {'norme_de_reference': 'F',
                                       'code': 'G', 'article': 'H'}),
    'Programme de rattachement': (5, 'D', {'programme_code': 'F',
                                           'programme_libelle': 'G',
                                           'programme_titre': 'H'}),
}
DF_NUMERIQUES = {'realisation_2024', 'prevision_2025', 'prevision_2026',
                 'beneficiaire_nombre'}


def enrichir_depenses_fiscales(chemin, entrees):
    """Ajoute au socle de chaque dépense fiscale ce que les autres feuilles disent."""
    par_num = {e['socle']['numero']: e for e in entrees}
    wb = openpyxl.load_workbook(chemin, data_only=True)
    vus = {nom: 0 for nom in DF_FEUILLES}
    for nom, (debut, col_num, colonnes) in DF_FEUILLES.items():
        if nom not in wb.sheetnames:
            continue
        ws = wb[nom]
        for i in range(debut, ws.max_row + 1):
            num = _txt(ws[f'{col_num}{i}'].value)
            e = par_num.get(num)
            if e is None:
                continue
            vus[nom] += 1
            for cle, col in colonnes.items():
                v = ws[f'{col}{i}'].value
                if cle in DF_NUMERIQUES:
                    n = _num(v)
                    e['socle'][cle] = n
                    if n is None and _txt(v):
                        e['socle'][cle + '_mention'] = _txt(v)
                else:
                    e['socle'][cle] = _txt(v)
    return vus


# ----------------------------------------------- l'exécution du budget, RAP 2024
# Deux feuilles du classeur des dépenses, à la maille mission, programme, action.
# Elles portent l'**exécuté** là où le PAP porte le **prévu** : c'est ce qui
# permet de dire si une prévision de crédit se réalise.
RAP_CREDITS_COLONNES = {
    'nature_ligne': 'A', 'budget': 'B', 'code': 'C', 'libelle': 'D',
    'ae_titre2_2023': 'E', 'ae_hors_titre2_2023': 'F', 'ae_2023': 'G',
    'cp_titre2_2023': 'H', 'cp_hors_titre2_2023': 'I', 'cp_2023': 'J',
    'ae_titre2_2024_lfi': 'K', 'ae_hors_titre2_2024_lfi': 'L',
    'ae_2024_lfi': 'M', 'cp_titre2_2024_lfi': 'N',
    'cp_hors_titre2_2024_lfi': 'O', 'cp_2024_lfi': 'P',
    'ae_2024_execute': 'Q', 'cp_2024_execute': 'T',
}
RAP_EMPLOIS_COLONNES = {
    'nature_ligne': 'A', 'budget': 'B', 'code': 'C', 'libelle': 'D',
    'realisation_2023': 'F', 'lfi_lfr_2024': 'G', 'realisation_2024': 'I',
    'ecart_a_lfi_2024': 'J', 'mesures_perimetre_2024': 'K',
    'mesures_transfert_2024': 'L',
}


def _feuille_rap(chemin, onglet, colonnes, prefixe):
    wb = openpyxl.load_workbook(chemin, data_only=True)
    if onglet not in wb.sheetnames:
        return []
    ws = wb[onglet]
    out = []
    for i in range(4, ws.max_row + 1):
        v = {k: ws[f'{c}{i}'].value for k, c in colonnes.items()}
        if not _txt(v['libelle']):
            continue
        socle = {}
        for k in colonnes:
            socle[k] = (_txt(v[k]) if k in ('nature_ligne', 'budget', 'code',
                                            'libelle') else _num(v[k]))
        out.append({'id': f"{prefixe}-{_txt(v['code']) or i}-{i}",
                    'ligne': i, 'socle': socle})
    return out


def des_credits_rap(chemin):
    return _feuille_rap(chemin, 'Crédits RAP 2024', RAP_CREDITS_COLONNES, 'RAPC')


def des_emplois_rap(chemin):
    return _feuille_rap(chemin, 'Emplois RAP 2024', RAP_EMPLOIS_COLONNES, 'RAPE')


# ---------------------------------------------------------------- les ODAC-ODAL
# Onglet « ODAC-ODAL » du classeur ETP et agences : les organismes divers
# d'administration centrale et locale, classés par fonction, avec le même jeu de
# régimes que les opérateurs et une colonne qui dit lesquels sont déjà traités
# au titre des opérateurs du PLF — celle-là évite le double compte.
ODAC_REGIMES = {'D': 'vente', 'E': 'suppression', 'F': 'epic_musee',
                'G': 'epic_etudes', 'H': 'internalisation'}


def des_odac(chemin, out_total=None):
    out_total = out_total if out_total is not None else {}
    ws = openpyxl.load_workbook(chemin, data_only=True)['ODAC-ODAL']
    out, fonction = [], None
    for i in range(4, ws.max_row + 1):
        nom = _txt(ws[f'A{i}'].value)
        if not nom:
            continue
        if nom.lower().startswith('fonction '):
            fonction = nom
            continue
        # La ligne 4 est le total de l'onglet — « ODAC 2023 Insee 2025 », 700
        # organismes — et non un organisme. Elle se garde à part.
        if nom.lower().startswith('odac '):
            out_total.update({'intitule': nom, 'nombre': _num(ws[f'B{i}'].value),
                              'deja_traites': _num(ws[f'C{i}'].value),
                              **{r: _num(ws[f'{c}{i}'].value)
                                 for c, r in ODAC_REGIMES.items()}})
            continue
        nb = _num(ws[f'B{i}'].value)
        if nb is None:
            continue
        regimes = [r for col, r in ODAC_REGIMES.items()
                   if _txt(ws[f'{col}{i}'].value)]
        out.append({
            'id': f"ODAC-{i}",
            'ligne': i,
            'socle': {'organisme': nom, 'fonction': fonction,
                      'nombre_entites': nb},
            'interpretation': {
                'deja_traite_en_operateur': bool(_txt(ws[f'C{i}'].value)),
                'regime': regimes[0] if len(regimes) == 1 else None,
                'regimes_multiples': regimes if len(regimes) > 1 else [],
                'sans_regime': not regimes,
            },
        })
    return out


# ------------------------------------------------- la nomenclature des missions
# Onglet « Annexe Etat » : le référentiel des missions et des programmes du
# budget de l'État. Il ne porte aucun montant — c'est la table qui permet de
# rattacher une ligne du PAP, une taxe affectée et un opérateur au même endroit.
def de_la_nomenclature(chemin):
    ws = openpyxl.load_workbook(chemin, data_only=True)['Annexe Etat']
    missions, programmes = {}, {}
    for i in range(2, ws.max_row + 1):
        code, libelle = _txt(ws[f'A{i}'].value), _txt(ws[f'B{i}'].value)
        if code and libelle:
            missions[code] = libelle
        genre = _txt(ws[f'G{i}'].value)
        if genre == 'P':
            num = _txt(ws[f'I{i}'].value)
            if num:
                programmes[num] = {'libelle': _txt(ws[f'J{i}'].value),
                                   'mission': _txt(ws[f'H{i}'].value)}
    return {'missions': missions, 'programmes': programmes}


# ------------------------------------------------------ l'arbre des économies
# Onglet « Détail Economies » du classeur de calculs. **Cette section ne vient
# pas du PLF.** Le classeur de calculs est entièrement de l'auteur : il n'y a
# pas ici un socle publié et une interprétation ajoutée, il n'y a qu'un calcul.
# Les entrées portent donc `classeur` — ce que la cellule dit, mot pour mot — et
# `lecture` — la place que ce module lui reconnaît dans l'arbre. Rien d'autre.
#
# L'arbre a trois étages. Deux totaux de tête (État, collectivités locales), des
# rubriques sous chacun, et sous quatre rubriques d'État un détail en « dont ».
# Une rubrique sans détail porte elle-même son code de destination.
#
#   D  restitué en CSG dès l'année 1        E  solde restitué ensuite
#   F  total supprimé — c'est lui l'économie valorisable
#   A  code de destination : où l'économie retombe une fois reclassée
#   G  l'hypothèse de chiffrage, en toutes lettres
#   I  une note de dérivation (base d'ETP, économie d'année 1)
#   L:P l'incidence par population, quand elle est allouée
#
# **Le code de la colonne A n'est pas la rubrique.** Une économie logée sous
# « chèques aux ménages » peut porter le code AE et rejoindre les aides à
# l'emploi. C'est cette migration que la colonne A trace, et la perdre revient
# à ne plus savoir qui supporte quoi.
ECO_DEBUT, ECO_FIN = 3, 43
ECO_TETES = {3: 'etat', 39: 'collectivites_locales'}
ECO_INCIDENCE = {'menages': 'L', 'actifs': 'M', 'retraites': 'N',
                 'entreprises': 'O', 'autres': 'P'}
# Les codes de destination, tels que le classeur les abrège. La table est écrite
# à la main : elle ne se devine pas des lettres.
ECO_CODES = {
    'AE': "aides à l'emploi et à l'apprentissage",
    'O': 'opérateurs',
    'E': 'aides aux entreprises',
    'M': 'chèques aux ménages',
    'AME': 'aide médicale de l’État',
    'AVP': 'audiovisuel public',
    'HU': "hébergement d'urgence",
    'APD': 'aide publique au développement',
    'A': 'subventions aux associations',
    'CG': 'charges courantes et achats',
    'MS': 'masse salariale',
}


def des_economies(chemin):
    ws = openpyxl.load_workbook(chemin, data_only=True)['Détail Economies']
    out, tete, rubrique = [], None, None
    for i in range(ECO_DEBUT, ECO_FIN + 1):
        intitule = _txt(ws[f'C{i}'].value)
        if not intitule:
            continue
        detail = intitule.lower().startswith('dont') or intitule == 'autres'
        if i in ECO_TETES:
            niveau, tete, rubrique = 'total', ECO_TETES[i], None
        elif detail:
            niveau = 'detail'
        else:
            niveau, rubrique = 'rubrique', intitule
        code = _txt(ws[f'A{i}'].value)
        incidence = {k: _num(ws[f'{c}{i}'].value)
                     for k, c in ECO_INCIDENCE.items()}
        out.append({
            'id': f'ECO-{i}',
            'ligne': i,
            'classeur': {
                'intitule': intitule,
                'restitue_annee_1_md_eur': _num(ws[f'D{i}'].value),
                'solde_restitue_ensuite_md_eur': _num(ws[f'E{i}'].value),
                'total_supprime_md_eur': _num(ws[f'F{i}'].value),
                'hypothese': _txt(ws[f'G{i}'].value),
                'note': _txt(ws[f'I{i}'].value),
                'code_destination': code,
                'incidence_md_eur': {k: v for k, v in incidence.items()
                                     if v is not None},
            },
            'lecture': {
                'niveau': niveau,
                'perimetre': tete,
                'rubrique': rubrique if niveau == 'detail' else None,
                'destination': ECO_CODES.get(code) if code else None,
                'code_inconnu': bool(code and code not in ECO_CODES),
            },
        })
    return out


# ------------------------------------------------------- la grande synthèse
# Onglet « Gages » du classeur de calculs : les douze lignes qui font le total
# supprimé à terme, chacune avec sa source citée. C'est la seule pièce du corpus
# qui écrive, à côté de chaque montant, d'où il vient. Elle est reprise telle
# quelle — ni recomposée, ni arrondie.
GAGES_LIGNES = range(8, 27)
GAGES_COLONNES = {
    'intitule': 'B', 'gain_direct_md_eur': 'C', 'effet': 'D',
    'gain_indirect_md_eur': 'E', 'commentaire_indirect': 'F',
    'total_supprime_md_eur': 'G', 'gisement_en_sus_md_eur': 'H',
    'commentaire_gisement': 'I', 'source': 'J',
}


def de_la_grande_synthese(chemin):
    ws = openpyxl.load_workbook(chemin, data_only=True)['Gages']
    out = []
    for i in GAGES_LIGNES:
        v = {k: ws[f'{c}{i}'].value for k, c in GAGES_COLONNES.items()}
        if not _txt(v['intitule']):
            continue
        out.append({
            'id': f'GAGE-{i}',
            'ligne': i,
            'classeur': {
                'intitule': _txt(v['intitule']),
                'gain_direct_md_eur': _num(v['gain_direct_md_eur']),
                'effet': _txt(v['effet']),
                'gain_indirect_md_eur': _num(v['gain_indirect_md_eur']),
                'commentaire_indirect': _txt(v['commentaire_indirect']),
                'total_supprime_md_eur': _num(v['total_supprime_md_eur']),
                'gisement_en_sus_md_eur': _num(v['gisement_en_sus_md_eur']),
                'commentaire_gisement': _txt(v['commentaire_gisement']),
                'source': _txt(v['source']),
            },
        })
    return out


# ------------------------------------------------------------------ la génération
def generer(dst, taxes=None, operateurs=None, depenses_fiscales=None,
            depenses_bg=None, calculs=None):
    if openpyxl is None:
        print('openpyxl absent : pip install openpyxl --break-system-packages')
        return 2
    socle = {
        '_revision': {
            'version': 'socle budgétaire v1',
            'exercice': 'PLF 2026',
            'objet': "lecture formalisée des classeurs sectoriels — le socle "
                     "que le PLF publie, et l'interprétation que l'auteur y a "
                     "ajoutée, tenus séparés",
            'regle': "Les classeurs de l'auteur restent la source officielle. "
                     "Rien n'est corrigé ici, rien n'est inventé : une cellule "
                     "vide reste vide.",
            'produit_par': 'appareil/socle_budgetaire.py',
            'grille': 'methode/grille_lecture_budgetaire.md',
        },
    }
    comptes = {}
    if taxes:
        socle['taxes_affectees'] = des_taxes(taxes)
        comptes['taxes_affectees'] = len(socle['taxes_affectees'])
    if operateurs:
        socle['operateurs'] = des_operateurs(operateurs)
        socle['agences'] = des_agences(operateurs)
        odac_total = {}
        socle['odac_odal'] = des_odac(operateurs, odac_total)
        socle['odac_total'] = odac_total
        socle['nomenclature'] = de_la_nomenclature(operateurs)
        comptes['operateurs'] = len(socle['operateurs'])
        comptes['odac_odal'] = len(socle['odac_odal'])
        comptes['programmes'] = len(socle['nomenclature']['programmes'])
    if depenses_fiscales:
        socle['depenses_fiscales'] = des_depenses_fiscales(depenses_fiscales)
        vus = enrichir_depenses_fiscales(depenses_fiscales,
                                         socle['depenses_fiscales'])
        socle['_revision']['feuilles_df_reprises'] = vus
        comptes['depenses_fiscales'] = len(socle['depenses_fiscales'])
    if depenses_bg:
        socle['pap'] = du_pap(depenses_bg)
        socle['bg_synthese'] = du_bg_synthese(depenses_bg)
        socle['credits_rap_2024'] = des_credits_rap(depenses_bg)
        socle['emplois_rap_2024'] = des_emplois_rap(depenses_bg)
        comptes['pap'] = len(socle['pap'])
        comptes['credits_rap'] = len(socle['credits_rap_2024'])
        comptes['emplois_rap'] = len(socle['emplois_rap_2024'])
    if calculs:
        socle['economies'] = des_economies(calculs)
        socle['grande_synthese'] = de_la_grande_synthese(calculs)
        comptes['economies'] = len(socle['economies'])
        comptes['grande_synthese'] = len(socle['grande_synthese'])
    socle['comptes'] = comptes

    with open(dst, 'w', encoding='utf-8') as f:
        json.dump(socle, f, ensure_ascii=False, indent=1)
        f.write('\n')
    print(f"{dst} écrit — "
          + ' · '.join(f'{k} {v}' for k, v in sorted(comptes.items())))
    return 0


if __name__ == '__main__':
    args, opts = sys.argv[1:], {}
    if not args:
        print(__doc__)
        sys.exit(2)
    dst, reste = args[0], args[1:]
    while reste:
        cle = reste.pop(0).lstrip('-').replace('-', '_')
        opts[cle] = reste.pop(0) if reste else None
    sys.exit(generer(dst, opts.get('taxes'), opts.get('operateurs'),
                     opts.get('depenses_fiscales'), opts.get('depenses_bg'),
                     opts.get('calculs')))
