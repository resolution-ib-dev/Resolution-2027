#!/usr/bin/env python3
"""Les tablettes des prélèvements et du produit privé.

Inscrit le 20261020. Écrit, depuis les référentiels et les trois calculs
(`rapprocher_7classements.py`, `portage_prelevements.py`,
`produit_prive_restant.py`), une tablette par question, en deux formes :

  livrables/tablettes/T<n>_<nom>.tsv      lisible sans outil, versable
  livrables/tablettes/Tablettes_prelevements.xlsx   mise en forme, un onglet par
                                          tablette, et un onglet « Sources »

Rien n'est publié : le site se tient à la main (Site-ETNP), après validation.
Montants en milliards d'euros courants sauf mention.
"""
import contextlib
import csv
import io
import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import portage_prelevements as portage          # noqa: E402
import produit_prive_restant as ppr            # noqa: E402

RACINE = Path(__file__).resolve().parent.parent
REF = RACINE / "referentiels"
SORTIE = RACINE / "livrables/tablettes"


def tsv(nom):
    return list(csv.DictReader((REF / nom).open(encoding="utf-8"), delimiter="\t"))


def t3217():
    L = {}
    for r in tsv("insee_3217_payeur.tsv"):
        L[int(r["ligne"])] = {int(k[1:]): float(v) for k, v in r.items()
                              if k.startswith("y") and v not in ("", "n.d.")}
    return L


# --- Les tablettes. Chacune rend (titre, note, en-têtes, lignes, formats).

def t1_heure(comptes, heures):
    """Une heure de travail : ce qui est produit, pris, emprunté, laissé."""
    L = t3217()
    s = ppr.serie(comptes, "credits")
    a_ = ppr.serie(comptes, "auteur")
    lignes = []
    for a in sorted(heures):
        x, h = s[a], heures[a]
        prelev = (L[104][a] + L[105][a]) - x["cot_pub"] - x["imp_pub"] - comptes[ppr.D12][a]
        lignes.append([a, x["pib"] / h, x["vapub"] / h, x["pp"] / h, prelev / h,
                       (x["rg"] - prelev) / h, x["deficit"] / h, x["ppr"] / h, a_[a]["ppr"] / h])
    return ("Une heure de travail",
            "Euros courants par heure travaillée (emploi intérieur, Eurostat). Prélèvements : impôts et "
            "cotisations effectives, non recouvré retranché, crédits d'impôt non retranchés, hors cotisations "
            "et impôts sur la masse salariale publique. Variante « crédits » du calcul du produit privé restant ; "
            "dernière colonne, le classeur de l'auteur tel quel.",
            ["Année", "PIB", "dont production publique comptée à son coût", "Produit privé",
             "Prélevé (impôts et cotisations)", "Autres recettes publiques", "Emprunté (impôt futur)",
             "Produit privé restant", "Restant, classeur de l'auteur"],
            lignes, ["0"] + ["0.0"] * 8)


def t2_ppr(comptes):
    s = ppr.serie(comptes, "credits")
    a_ = ppr.serie(comptes, "auteur")
    lignes = []
    for a in ppr.ANNEES:
        x = s[a]
        g_pp, g_ppr = ppr.croissance(s, "pp", a), ppr.croissance(s, "ppr", a)
        lignes.append([a, x["pib"], x["vapub"], x["pp"], x["dg"], x["rg"], x["ppr"],
                       x["ppr"] / x["pp"], x["ppr"] / x["pib"], g_pp, g_ppr, a_[a]["ppr"]])
    return ("Produit privé restant, 1959-2025",
            "Produit privé = PIB − valeur ajoutée publique. Produit privé restant = produit privé − le plus grand "
            "des deux : dépense ou recette publique, hors flux internes au circuit public et hors production "
            "immobilisée pour compte propre. Crédits d'impôt gardés des deux côtés (arbitrage du 20261020). "
            "Croissances en volume, déflateur du PIB.",
            ["Année", "PIB", "Valeur ajoutée publique", "Produit privé", "Dépense publique",
             "Recette publique", "Produit privé restant", "en % du produit privé", "en % du PIB",
             "Croissance du produit privé", "Croissance du produit privé restant",
             "Restant, classeur de l'auteur"],
            lignes, ["0"] + ["#,##0.0"] * 6 + ["0.0%"] * 4 + ["#,##0.0"])


def t3_familles(comptes):
    L = t3217()
    lignes = []
    for a in range(1995, 2026):
        pib = comptes[ppr.PIB][a]
        csg = sum(L[n].get(a, 0) for n in (71, 72, 73))
        brut = L[104][a] + L[105][a] + L[106][a]
        lignes.append([a, L[99][a] / pib, csg / pib, L[75][a] / pib, L[78][a] / pib, L[5][a] / pib,
                       brut / pib, L[110][a] / pib])
    return ("La glissade de la cotisation à la contribution, 1995-2025",
            "En % du PIB. INSEE, tableau 3.217, édition du classeur de l'auteur ; PIB du tableau 1.101. "
            "« Prélèvements, notion du projet » : hors cotisations imputées et non recouvré, crédits d'impôt non "
            "retranchés. « Prélèvements obligatoires INSEE » : crédits d'impôt retranchés.",
            ["Année", "Cotisations sociales effectives", "CSG, CRDS et prélèvements sociaux",
             "Impôt sur le revenu, avant crédits", "Impôt sur les sociétés, avant crédits", "TVA",
             "Prélèvements, notion du projet", "Prélèvements obligatoires INSEE"],
            lignes, ["0"] + ["0.0%"] * 7)


def t4_passage(comptes):
    L = t3217()
    lignes = []
    for libelle, n, signe in (("Ensemble des impôts et cotisations sociales", 104, 1),
                              ("− non recouvrés", 105, 1), ("− cotisations imputées", 106, 1)):
        lignes.append([libelle] + [L[n][a] for a in (2023, 2024, 2025)])
    brut = [sum(L[n][a] for n in (104, 105, 106)) for a in (2023, 2024, 2025)]
    lignes.append(["= prélèvements, notion du projet (Eurostat de même)"] + brut)
    lignes.append(["− crédits d'impôt (convention INSEE)"] + [L[107][a] for a in (2023, 2024, 2025)])
    lignes.append(["= prélèvements obligatoires INSEE"] + [L[110][a] for a in (2023, 2024, 2025)])
    lignes.append(["PIB"] + [comptes[ppr.PIB][a] for a in (2023, 2024, 2025)])
    lignes.append(["Notion du projet, en % du PIB"] + [b / comptes[ppr.PIB][a] for b, a in zip(brut, (2023, 2024, 2025))])
    lignes.append(["Prélèvements obligatoires INSEE, en % du PIB"] + [L[110][a] / comptes[ppr.PIB][a] for a in (2023, 2024, 2025)])
    return ("Du brut aux prélèvements obligatoires",
            "INSEE, tableau 3.217, édition du classeur de l'auteur. La comparaison avant / après de la doctrine "
            "reste sur la notion INSEE (arbitrage du 20261020).",
            ["", "2023", "2024", "2025"], lignes, ["@"] + ["#,##0.0"] * 3, {7: ["@"] + ["0.0%"] * 3, 8: ["@"] + ["0.0%"] * 3})


def t5_portage():
    with contextlib.redirect_stdout(io.StringIO()):
        portage.main()
    R = portage.RESULTATS
    lignes = []
    for nom, cle in (("Calcul de l'auteur", "auteur"), ("Corrigé sur pièces", "corrige")):
        r = R[cle]
        lignes.append([nom, r["base"], r["entreprises"], r["menages"], r["entreprises"] / r["base"],
                       r["entreprises_hors_pas"] / r["base"]])
    return ("Qui remet le prélèvement privé, 2025",
            "Part remise au fisc par les entreprises, avec et sans le prélèvement à la source (ratio de l'auteur, "
            "sans source). Corrections : crédits d'impôt non retranchés, CSG du capital gardée, flux publics retirés "
            "sur pièces Eurostat. Détail : livrables/portage_prelevements_20261020.md.",
            ["", "Base privée", "Remis par les entreprises", "Remis par les ménages",
             "Part des entreprises", "Part des entreprises, hors prélèvement à la source"],
            lignes, ["@"] + ["#,##0.0"] * 3 + ["0.0%"] * 2)


def t6_taxud():
    t = tsv("prelevements_7classements_1995_2024.tsv")
    lignes = []
    for a in range(1995, 2025):
        f = defaultdict(float)
        for r in t:
            f[r["tx"]] += int(r[f"y{a}"]) / 1000
        tot = f["Travail"] + f["Consommation"] + f["Capital"]
        lignes.append([a, f["Travail"] / tot, f["Consommation"] / tot, f["Capital"] / tot, tot])
    return ("Sur quoi porte le prélèvement, 1995-2024",
            "Fonction économique de la Commission européenne (TAXUD), table du fiscaliste, hors TEOM et "
            "cotisations imputées. Aucune catégorie « entreprise » : trois bases.",
            ["Année", "Travail", "Consommation", "Capital", "Total (Md€)"],
            lignes, ["0", "0.0%", "0.0%", "0.0%", "#,##0.0"])


def t7_classements():
    t = tsv("prelevements_7classements_1995_2024.tsv")
    noms = {"qj": "Qualification juridique", "ins": "INSEE", "sec": "Eurostat (SEC 2010)", "fmi": "FMI",
            "ocde": "OCDE", "grd": "ONU (UNU-WIDER)", "tx": "TAXUD"}
    lignes = []
    for k, nom in noms.items():
        f = defaultdict(float)
        for r in t:
            f[r[k]] += int(r["y2024"]) / 1000
        for cat, m in sorted(f.items(), key=lambda x: -x[1]):
            lignes.append([nom, cat, m])
    return ("Un même euro, sept classements, 2024",
            "Prélèvements bruts 2024, 1 335,5 Md€ dont TEOM et cotisations imputées. Table du fiscaliste.",
            ["Classement", "Catégorie", "Md€"], lignes, ["@", "@", "#,##0.0"])


def ecrire(tablettes):
    SORTIE.mkdir(parents=True, exist_ok=True)
    for i, tb in enumerate(tablettes, 1):
        titre, note, entetes, lignes = tb[:4]
        nom = f"T{i}_" + "".join(c if c.isalnum() else "_" for c in titre.lower())[:40].strip("_")
        with (SORTIE / f"{nom}.tsv").open("w", encoding="utf-8") as o:
            o.write("\t".join(entetes) + "\n")
            for l in lignes:
                o.write("\t".join("" if x is None else (f"{x:.6g}" if isinstance(x, float) else str(x))
                                  for x in l) + "\n")
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    gras, gris = Font(bold=True), PatternFill("solid", fgColor="EEEEEE")
    for i, tb in enumerate(tablettes, 1):
        titre, note, entetes, lignes, formats = tb[:5]
        speciaux = tb[5] if len(tb) > 5 else {}
        ws = wb.create_sheet(f"T{i}")
        ws["A1"], ws["A1"].font = f"T{i} — {titre}", Font(bold=True, size=13)
        ws["A2"] = note
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(len(entetes), 4))
        ws.row_dimensions[2].height = 75
        for j, h in enumerate(entetes, 1):
            c = ws.cell(4, j, h)
            c.font, c.fill = gras, gris
            c.alignment = Alignment(wrap_text=True, vertical="bottom")
        ws.row_dimensions[4].height = 45
        for k, l in enumerate(lignes):
            fm = speciaux.get(k, formats)
            for j, x in enumerate(l, 1):
                c = ws.cell(5 + k, j, x)
                c.number_format = fm[j - 1]
        for j in range(1, len(entetes) + 1):
            ws.column_dimensions[get_column_letter(j)].width = 44 if j == 1 and formats[0] == "@" else 16
        ws.freeze_panes = "B5"
    src = wb.create_sheet("Sources")
    for k, l in enumerate(SOURCES, 1):
        src.cell(k, 1, l)
    src.column_dimensions["A"].width = 140
    wb.save(SORTIE / "Tablettes_prelevements.xlsx")


SOURCES = [
    "Sources",
    "INSEE, comptes nationaux annuels base 2020 : tableaux 1.101, 1.103, 3.201, 3.217, 7.301, 7.401 "
    "(classeurs de l'auteur T_3217_Impots_0902 et T_1101_1103_PIBtoPPR).",
    "Eurostat : gov_10a_main, gov_10a_taxag, nasa_10_nf_tr, nama_10_a10_e (extraction du 20 octobre 2026).",
    "Table « sept classements » d'un fiscaliste : INSEE 3.217, liste nationale TAXUD, OCDE Revenue Statistics, "
    "FMI GFSM 2014, UNU-WIDER GRD, 1995-2024.",
    "Méthode du produit privé restant : M. N. Rothbard, Power and Market, chapitre 5.",
    "Calculs : appareil/produit_prive_restant.py, portage_prelevements.py, rapprocher_7classements.py.",
]


def main():
    comptes = ppr.charger()
    heures = ppr.heures()
    tablettes = [t1_heure(comptes, heures), t2_ppr(comptes), t3_familles(comptes), t4_passage(comptes),
                 t5_portage(), t6_taxud(), t7_classements()]
    ecrire(tablettes)
    for i, tb in enumerate(tablettes, 1):
        print(f"T{i} {tb[0]} : {len(tb[3])} lignes")
    return 0


if __name__ == "__main__":
    sys.exit(main())
